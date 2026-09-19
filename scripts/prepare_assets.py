from pathlib import Path
import json, shutil
from openpyxl import load_workbook
from pptx import Presentation

src = Path(r"C:\Users\meera\AppData\Local\Temp")
root = Path(__file__).resolve().parents[1]
assets = root / "public" / "assets"
assets.mkdir(parents=True, exist_ok=True)

for name in ["AI Skilling Module 1.pptx", "AI Skilling Module 2.pptx", "AI Skilling Module 3.pptx", "AI Skilling - Combined.pptx", "AI SKILLING CONCEPT NOTE.docx"]:
    shutil.copy2(src / name, assets / name)

wb = load_workbook(src / "AI Skilling Project Planning.xlsx", data_only=True, read_only=True)

def clean(v):
    if v is None: return ""
    if hasattr(v, "isoformat"): return v.isoformat()
    return str(v).strip()

ws = wb["Session Details"]
headers = [clean(c.value) for c in ws[1]]
sessions = []
for row in ws.iter_rows(min_row=2, values_only=True):
    if not any(v is not None and str(v).strip() for v in row): continue
    item = {headers[i]: clean(row[i]) for i in range(min(len(headers), len(row)))}
    item['Participant Type'] = item.get('Participant Type (Student/ Teacher/ Headmaster/ Others)', '')
    item['Number participants'] = item.get('Number of participants', '')
    item['Overall Session observation'] = item.get('Over all Session observation', '')
    sessions.append(item)

wa = wb["Teacher Attendance Sheet"]
aheaders = [clean(c.value) for c in wa[2]][:11]
attendance = []
for row in wa.iter_rows(min_row=3, values_only=True):
    vals = row[:11]
    if not any(v is not None and str(v).strip() for v in vals): continue
    item = {aheaders[i]: clean(vals[i]) for i in range(len(aheaders))}
    item['Participant Type'] = item.get('Participant Type (Student/ Teacher/ Headmaster/ Others)', '')
    item['Name'] = item.get('Name of the participant', '')
    item['Pre-assessment Score'] = item.get('Pre-Assessment Score', '')
    item['Post-assessment Score'] = item.get('Post ASsessment Score', '')
    item.pop('Phone number', None)
    item['Name'] = item.get('Name of the participant', '')
    attendance.append(item)

(root / "app" / "data.ts").write_text("export const sessions = " + json.dumps(sessions, ensure_ascii=False) + " as const;\nexport const attendance = " + json.dumps(attendance, ensure_ascii=False) + " as const;\n", encoding="utf-8")

# Extract the visual marks used in the official decks for the site header.
prs = Presentation(src / "AI Skilling Module 1.pptx")
seen = 0
for slide in list(prs.slides)[:3]:
    for shape in slide.shapes:
        if getattr(shape, "shape_type", None) == 13:
            try:
                image = shape.image
            except ValueError:
                continue
            seen += 1
            (assets / f"deck-mark-{seen}.{image.ext}").write_bytes(image.blob)
